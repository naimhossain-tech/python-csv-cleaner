import argparse
import csv
import sys
from pathlib import Path

import openpyxl


def clean_value(value):
    return value.strip() if isinstance(value, str) else value


def validate_headers(raw_headers):
    if not raw_headers or all(value is None or str(value).strip() == "" for value in raw_headers):
        raise ValueError("Header row is missing.")

    headers = [str(value).strip() if value is not None else "" for value in raw_headers]

    if any(not header for header in headers):
        raise ValueError("Empty column names are not allowed.")

    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column names after trimming.")

    return headers


def clean_csv(source):
    with source.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        raw_headers = next(reader, None)
        headers = validate_headers(raw_headers)
        rows = []
        seen = set()
        total = 0
        blank_removed = 0
        duplicate_removed = 0

        for values in reader:
            total += 1
            if len(values) > len(headers):
                raise ValueError(f"Too many values on CSV line {reader.line_num}.")
            values += [""] * (len(headers) - len(values))
            cleaned = [value.strip() for value in values]

            if all(value == "" for value in cleaned):
                blank_removed += 1
                continue

            key = tuple(cleaned)
            if key in seen:
                duplicate_removed += 1
                continue

            seen.add(key)
            rows.append(cleaned)

    return headers, rows, total, blank_removed, duplicate_removed


def clean_excel(source):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        raw_headers = next(iterator, None)
        headers = validate_headers(raw_headers)
        rows = []
        seen = set()
        total = 0
        blank_removed = 0
        duplicate_removed = 0

        for values in iterator:
            total += 1
            cleaned = [clean_value(value) for value in values[:len(headers)]]
            cleaned += [None] * (len(headers) - len(cleaned))

            if all(value is None or value == "" for value in cleaned):
                blank_removed += 1
                continue

            key = tuple(cleaned)
            if key in seen:
                duplicate_removed += 1
                continue

            seen.add(key)
            rows.append(cleaned)

        return headers, rows, total, blank_removed, duplicate_removed
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(
        description="Clean CSV or Excel files and remove blank or duplicate rows."
    )
    parser.add_argument("input", help="Path to an input .csv or .xlsx file")
    parser.add_argument(
        "-o", "--output",
        help="Output file path (default: cleaned_output with the input extension)"
    )
    args = parser.parse_args()

    source = Path(args.input)
    extension = source.suffix.lower()

    if extension not in (".csv", ".xlsx"):
        print("Error: Input must be a .csv or .xlsx file.")
        return 1

    output = Path(args.output) if args.output else Path(f"cleaned_output{extension}")

    if output.suffix.lower() != extension:
        print(f"Error: Output file must also use the {extension} extension.")
        return 1

    if not source.is_file():
        print(f"Error: Input file not found: {source}")
        return 1

    try:
        if source.resolve() == output.resolve():
            print("Error: Input and output files must be different.")
            return 1

        if extension == ".csv":
            headers, rows, total, blank_removed, duplicate_removed = clean_csv(source)
        else:
            headers, rows, total, blank_removed, duplicate_removed = clean_excel(source)

        output.parent.mkdir(parents=True, exist_ok=True)

        if extension == ".csv":
            with output.open("w", newline="", encoding="utf-8") as file:
                writer = csv.writer(file)
                writer.writerow(headers)
                writer.writerows(rows)
        else:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            workbook.save(output)
            workbook.close()

        removed = blank_removed + duplicate_removed
        kept = len(rows)
        removal_rate = (removed / total * 100) if total else 0.0

        print("\n" + "=" * 36)
        print("         FILE CLEANING REPORT")
        print("=" * 36)
        print("Status:                SUCCESS")
        print(f"Input file:            {source}")
        print(f"Output file:           {output}")
        print("-" * 36)
        print(f"Total rows checked:    {total}")
        print(f"Blank rows removed:    {blank_removed}")
        print(f"Duplicate rows removed:{duplicate_removed:>5}")
        print(f"Total rows removed:    {removed}")
        print(f"Rows kept:             {kept}")
        print(f"Removal rate:          {removal_rate:.2f}%")
        print("=" * 36)
        return 0

    except (OSError, UnicodeError, csv.Error, ValueError, openpyxl.utils.exceptions.InvalidFileException) as error:
        print(f"Error: {error}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
